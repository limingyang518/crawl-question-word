"""导出极光题库中当前登录账号可访问的题目为 Word 文档。

运行：python crawl_question.py
先在网站重新登录，运行时粘贴新的 token（输入不会显示）。
也可通过环境变量 JIGUANG_TOKEN 提供 token。
"""
import argparse
import base64
import getpass
from html.parser import HTMLParser
import json
import os
from pathlib import Path
import re
import sys
import time
from urllib.parse import quote, urlsplit

import requests

BASE_URL = "https://api.ji-guang.top/api/v1"
TOKEN = ""  # 可在这里填入重新登录后的 token，建议使用环境变量或运行时输入。
SAVE_DIR = Path.home() / "Desktop" / "题库"


class CrawlError(Exception):
    pass


class AuthError(CrawlError):
    pass


def normalize_token(token):
    token = token.strip()
    if token.lower().startswith("bearer "):
        token = token[7:].strip()
    if not token:
        raise AuthError("未提供登录 token。请先在网站登录，再输入新的 token。")
    # 仅检查本地过期时间以帮助诊断；不代替服务器的鉴权。
    try:
        payload = json.loads(base64.urlsafe_b64decode(token.split(".")[1] + "==="))
        if float(payload.get("exp", float("inf"))) <= time.time():
            raise AuthError("登录 token 已过期，请在网站重新登录后换用新的 token。")
    except (ValueError, IndexError, TypeError, UnicodeError):
        pass
    return token


def safe_filename(name):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name)).strip().rstrip(". ")
    name = name[:90] or "未命名"
    if name.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *[f"COM{i}" for i in range(1, 10)], *[f"LPT{i}" for i in range(1, 10)]}:
        name = "_" + name
    return name


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


class _ContentParser(HTMLParser):
    """保留 HTML 的文字和段落分隔，不执行或下载网页内容。"""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1
        if self.hidden:
            return
        if tag in ("br", "p", "div", "li", "tr", "h1", "h2", "h3"):
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" ")
        elif tag == "img":
            attrs = dict(attrs)
            self.parts.append("[图片：" + (attrs.get("alt") or attrs.get("src") or "未提供地址") + "]")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)
        elif not self.hidden and tag in ("p", "div", "li", "tr", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def readable_text(value):
    if isinstance(value, list):
        return "\n".join(readable_text(item) for item in value)
    if isinstance(value, dict):
        return "\n".join(f"{key}：{readable_text(item)}" for key, item in value.items())
    parser = _ContentParser()
    parser.feed(as_text(value))
    parser.close()
    text = "".join(parser.parts).replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
    # Word XML 不允许控制字符或代理字符，过滤以免整个章节保存失败。
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def save_word(rows, path, subject, chapter, include_explanation=True, failures=0):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.3)
    for name, size, font in (("Normal", 11, "宋体"), ("Title", 22, "黑体"), ("Heading 1", 14, "黑体")):
        style = doc.styles[name]
        style.font.name, style.font.size = font, Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing, normal.space_after = 1.25, Pt(5)
    normal.widow_control = True
    heading = doc.styles["Heading 1"].paragraph_format
    heading.space_before, heading.space_after = Pt(12), Pt(5)
    heading.keep_with_next = True
    doc.core_properties.title = f"{subject} {chapter} 题库"
    doc.core_properties.author = ""
    doc.add_paragraph(f"{subject} {chapter} 题库", "Title")
    doc.add_paragraph(f"本章共收录 {len(rows)} 道题，按题目、选项和答案排列。" + ("有解析的题目附解析。" if include_explanation else ""))
    if failures:
        doc.add_paragraph(f"本文件为部分结果，爬取过程中有 {failures} 项失败。")
    for index, row in enumerate(rows, 1):
        doc.add_paragraph(f"第{index}题", "Heading 1")
        doc.add_paragraph(row["题目"] or "题干未提供")
        for key, value in row["选项"].items():
            if value:
                paragraph = doc.add_paragraph(f"{key}．{value}")
                paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph = doc.add_paragraph()
        paragraph.add_run("答案：").bold = True
        paragraph.add_run(row["答案"] or "未提供")
        if include_explanation and row["解析"]:
            paragraph = doc.add_paragraph()
            paragraph.add_run("解析：").bold = True
            paragraph.add_run(row["解析"])
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.docx")
    try:
        doc.save(temporary)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


class Client:
    def __init__(self, token, base_url=BASE_URL):
        parsed = urlsplit(base_url)
        if parsed.scheme != "https" or not parsed.netloc or parsed.fragment or parsed.query:
            raise CrawlError("API 地址必须是 HTTPS 接口根地址，不能包含 #/home 或查询参数。")
        self.base_url = base_url.rstrip("/")
        self.session = requests.Session()
        self.session.trust_env = False  # 保留原脚本禁用系统代理的设置。
        self.session.headers.update({
            "Authorization": "Bearer " + normalize_token(token),
            "Accept": "application/json",
            "User-Agent": "Mozilla/5.0",
        })

    def get(self, path, params=None, retries=3):
        url = self.base_url + path
        for attempt in range(retries):
            try:
                response = self.session.get(url, params=params, timeout=(10, 30))
            except requests.exceptions.SSLError as exc:
                raise CrawlError("HTTPS 证书校验失败，请检查系统时间和证书配置。") from exc
            except requests.RequestException as exc:
                if attempt + 1 == retries:
                    raise CrawlError(f"{path} 网络请求失败（{type(exc).__name__}）") from exc
                time.sleep(attempt + 1)
                continue
            if response.status_code == 401:
                raise AuthError("服务器返回 401：登录已失效，请重新登录并更新 token。")
            if response.status_code == 403:
                raise CrawlError(f"{path} 返回 403：当前账号没有访问权限。")
            if response.status_code == 429 or response.status_code >= 500:
                if attempt + 1 < retries:
                    time.sleep(attempt + 1)
                    continue
            if response.status_code != 200:
                raise CrawlError(f"{path} 请求失败，HTTP {response.status_code}")
            try:
                body = response.json()
            except ValueError as exc:
                raise CrawlError(f"{path} 返回了非 JSON 内容，请检查 API 地址。") from exc
            if not isinstance(body, dict):
                raise CrawlError(f"{path} 返回格式异常：顶层不是对象。")
            code = body.get("code")
            if str(code) == "401":
                raise AuthError("登录已失效，请重新登录并更新 token。")
            if body.get("success") is False or code not in (None, 0, 200, "0", "200"):
                raise CrawlError(f"{path} 接口错误：{as_text(body.get('message') or body.get('detail') or code)}")
            if "data" not in body or body["data"] is None:
                raise CrawlError(f"{path} 响应缺少 data。")
            return body["data"]
        raise CrawlError(f"{path} 请求失败。")

    def pages(self, path, field, params=None, page_size=100):
        skip = 0
        seen = set()
        while True:
            data = self.get(path, {**(params or {}), "skip": skip, "limit": page_size})
            if not isinstance(data, dict) or not isinstance(data.get(field), list):
                raise CrawlError(f"{path} 响应缺少列表字段 {field}。")
            items = data[field]
            if not items:
                if data.get("has_more") is True:
                    raise CrawlError(f"{path} 声称还有数据，但返回了空页。")
                return
            fresh = []
            for item in items:
                if not isinstance(item, dict):
                    raise CrawlError(f"{path} 列表项不是对象。")
                identity = item.get("id") or item.get("question_id")
                if not identity:
                    raise CrawlError(f"{path} 列表项缺少 id/question_id。")
                if str(identity) not in seen:
                    seen.add(str(identity))
                    fresh.append(item)
            if not fresh:
                raise CrawlError(f"{path} 分页重复，接口可能没有处理 skip；停止以避免漏题或死循环。")
            yield from fresh
            skip += len(items)
            if data.get("has_more") is False:
                return
            total = data.get("total")
            if isinstance(total, int) and skip >= total:
                return
            if data.get("has_more") is not True and total is None and len(items) < page_size:
                return


def question_row(question, subject, chapter, qid):
    if not isinstance(question, dict) or "content" not in question:
        raise CrawlError(f"题目 {qid} 详情缺少 content。")
    options = question.get("options") or {}
    if not isinstance(options, dict):
        raise CrawlError(f"题目 {qid} 的 options 格式异常，预期为对象。")
    return {
        "学科": subject, "章节": chapter, "题目ID": str(qid),
        "题目": readable_text(question.get("content")),
        "选项": {str(key): readable_text(value) for key, value in sorted(options.items(), key=lambda item: str(item[0]))},
        "答案": readable_text(question.get("correct_answer")),
        "解析": readable_text(question.get("ai_explanation")) or readable_text(question.get("explanation")),
        "难度": as_text(question.get("difficulty")), "题型": as_text(question.get("q_type")),
    }


def crawl(client, tag, save_dir, page_size=100, delay=0.15, include_explanation=True):
    from docx import Document  # 在请求之前检查 Word 写入依赖。

    print("正在获取学科列表...")
    subjects = list(client.pages("/user/subjects/", "items", {"tag": tag}, page_size))
    print(f"发现 {len(subjects)} 个学科")
    if not subjects:
        raise CrawlError(f"标签 {tag!r} 下没有学科，请检查 --tag。")
    failures = saved = question_count = 0
    for subject in subjects:
        subject_name = as_text(subject.get("name")) or str(subject["id"])
        sid = quote(str(subject["id"]), safe="")
        print(f"\n学科：{subject_name}")
        try:
            detail = client.get(f"/user/subjects/{sid}")
            if not isinstance(detail, dict) or not isinstance(detail.get("chapters"), list):
                raise CrawlError("学科详情缺少 chapters 列表。")
        except AuthError:
            raise
        except CrawlError as exc:
            failures += 1
            print(f"[失败] {exc}")
            continue
        for chapter in detail["chapters"]:
            rows = []
            chapter_failures = 0
            try:
                if not isinstance(chapter, dict) or not chapter.get("id"):
                    raise CrawlError("章节缺少 id。")
                chapter_name = as_text(chapter.get("name")) or str(chapter["id"])
                cid = quote(str(chapter["id"]), safe="")
                print(f"章节：{chapter_name}")
                for item in client.pages(f"/user/chapters/{cid}", "questions", {"mode": "practice"}, page_size):
                    qid = item.get("id") or item.get("question_id")
                    try:
                        question = client.get(f"/user/questions/{quote(str(qid), safe='')}")
                        rows.append(question_row(question, subject_name, chapter_name, qid))
                        print(f"[成功 {len(rows)}] {as_text(question.get('content'))[:40]}")
                    except AuthError:
                        raise
                    except CrawlError as exc:
                        chapter_failures += 1
                        print(f"[题目失败] {exc}")
                    time.sleep(delay)
            except AuthError:
                raise
            except CrawlError as exc:
                chapter_failures += 1
                print(f"[章节失败] {exc}")
            failures += chapter_failures
            if not rows:
                print("本章节没有成功获取的题目，不生成空 Word 文档。")
                continue
            # 完整 ID 防止重名学科/章节覆盖；部分结果明确标识。
            stem = safe_filename(subject_name) + "-" + safe_filename(chapter_name) + "-" + safe_filename(chapter['id'])
            if chapter_failures:
                stem += "-部分结果"
            path = Path(save_dir) / (stem + ".docx")
            try:
                save_word(rows, path, subject_name, chapter_name, include_explanation, chapter_failures)
            except (OSError, ValueError) as exc:
                failures += 1
                print(f"[保存失败] {path.name}：{exc}")
                continue
            saved += 1
            question_count += len(rows)
            print(f"已保存 {path}（{len(rows)} 道题，本章节失败 {chapter_failures}）")
    print(f"\n结果：保存 {saved} 个文件、{question_count} 道题；失败 {failures} 项。")
    if failures or not saved:
        raise CrawlError("爬取未全部成功，请根据上方提示处理后重试。")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--tag", default="三年级", help="学科标签（默认：三年级）")
    parser.add_argument("--save-dir", type=Path, default=SAVE_DIR, help="Word 保存目录")
    parser.add_argument("--page-size", type=int, default=100, help="每页题目数（1-10000）")
    parser.add_argument("--no-explanation", action="store_true", help="不导出解析，只保留题目、选项和答案")
    args = parser.parse_args(argv)
    if not 1 <= args.page_size <= 10000:
        parser.error("--page-size 必须在 1 到 10000 之间")
    client = None
    try:
        token = os.environ.get("JIGUANG_TOKEN", "").strip() or TOKEN.strip()
        if not token:
            if not sys.stdin.isatty():
                raise AuthError("请设置 JIGUANG_TOKEN 或在交互式终端运行并输入 token。")
            token = getpass.getpass("请输入重新登录后的 token（输入不显示）：")
        client = Client(token)
        crawl(client, args.tag, args.save_dir, args.page_size, include_explanation=not args.no_explanation)
        return 0
    except (CrawlError, ImportError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        if isinstance(exc, ImportError):
            print("请安装依赖：python -m pip install requests python-docx", file=sys.stderr)
        return 1
    except (KeyboardInterrupt, EOFError):
        print("已停止。", file=sys.stderr)
        return 130
    finally:
        if client is not None:
            client.session.close()


if __name__ == "__main__":
    sys.exit(main())
